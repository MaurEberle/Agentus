from __future__ import annotations

from pathlib import Path

from app.db import init
from app.db.bootstrap import load_bootstrap
from app.run.graph_models import AgentNetworkDocument, GraphNode
from app.run.knowledge import index_node, retrieve
from app.run.compile import compile_document
from app.run.controller import get_controller
from app.db.networks import NetworkRow, upsert_network
from app.db.engine import utc_now
from app.settings.models import AppSettingsPatch
from app.settings.service import patch_settings
from tests.run.conftest import mini_doc


def test_retrieve_before_complete_adds_context(monkeypatch, api_env) -> None:
    init()
    captured: list[str] = []

    def _complete(req):
        from app.runtime.models import CompletionResult

        captured.append(req.messages[0].content or "")
        return CompletionResult(content="hi", model=req.model, finish_reason="stop")

    monkeypatch.setattr(
        "app.runtime.completions.complete_live",
        lambda req, should_abort=None, on_progress=None: _complete(req),
    )
    monkeypatch.setattr(
        "app.run.harness.retrieve",
        lambda *a, **k: [
            __import__("app.run.knowledge", fromlist=["Snippet"]).Snippet(
                title="doc", section="s", text="secret snippet", score=1.0
            )
        ],
    )
    doc = mini_doc(startMessage="q")
    upsert_network(
        NetworkRow(
            id="net-1",
            name="mini",
            description=None,
            tags=[],
            document=doc,
            updated_at=utc_now(),
            last_used_at=None,
            last_run_id=None,
        )
    )
    patch_settings(AppSettingsPatch(active_network_id="net-1"))
    get_controller().start()
    thread = get_controller().thread
    if thread:
        thread.join(timeout=5)
    assert captured
    assert "secret snippet" in captured[0] or "No document context." in captured[0]


def test_ready_index_does_not_reread(monkeypatch, api_env, tmp_path) -> None:
    from app.runtime.models import EmbedResult

    init()
    calls = {"n": 0}

    def _embed(req):
        calls["n"] += 1
        return EmbedResult(vectors=[[0.1, 0.2] for _ in req.texts], dimension=2, model=req.model)

    monkeypatch.setattr("app.runtime.embeddings.embed", _embed)
    upsert_network(
        NetworkRow(
            id="net-1",
            name="kb",
            description=None,
            tags=[],
            document={"schemaVersion": 1, "name": "kb", "nodes": [], "edges": []},
            updated_at=utc_now(),
            last_used_at=None,
            last_run_id=None,
        )
    )
    folder = tmp_path / "notes"
    folder.mkdir()
    (folder / "note.md").write_text("# Hello\nworld", encoding="utf-8")
    node = GraphNode(
        id="k1",
        type="knowledge",
        position={"x": 0, "y": 0},
        data={"sourcePath": str(folder), "embeddingModel": "nomic-embed-text"},
    )
    data_dir = tmp_path / "appdata"
    assert index_node("net-1", node, data_dir=str(data_dir), force=True) == "ready"
    assert calls["n"] == 1

    def _read_forbidden(path):
        raise AssertionError(path)

    monkeypatch.setattr("app.run.knowledge._read", _read_forbidden)
    assert index_node("net-1", node, data_dir=str(data_dir)) == "ready"
    assert calls["n"] == 1


def test_index_uses_node_embedding_settings(monkeypatch, api_env) -> None:
    from app.runtime.models import EmbedRequest, EmbedResult

    init()
    captured: list[EmbedRequest] = []

    def _embed(req: EmbedRequest) -> EmbedResult:
        captured.append(req)
        return EmbedResult(
            vectors=[[0.1, 0.2] for _ in req.texts],
            dimension=2,
            model=req.model,
        )

    monkeypatch.setattr("app.runtime.embeddings.embed", _embed)
    upsert_network(
        NetworkRow(
            id="net-1",
            name="kb",
            description=None,
            tags=[],
            document={"schemaVersion": 1, "name": "kb", "nodes": [], "edges": []},
            updated_at=utc_now(),
            last_used_at=None,
            last_run_id=None,
        )
    )
    data_dir = load_bootstrap().data_dir
    folder = Path(api_env) / "kb"
    folder.mkdir(parents=True)
    (folder / "note.md").write_text("# Hello\nworld", encoding="utf-8")
    node = GraphNode(
        id="k1",
        type="knowledge",
        position={"x": 0, "y": 0},
        data={
            "sourcePath": str(folder),
            "embeddingProvider": "openai",
            "embeddingModel": "text-embedding-3-small",
            "embeddingCredentialId": "cred-1",
        },
    )
    assert index_node("net-1", node, data_dir=str(data_dir), force=True) == "ready"
    assert captured
    assert captured[0].model == "text-embedding-3-small"
    assert captured[0].provider == "openai"
    assert captured[0].credential_id == "cred-1"


def test_index_outside_data_dir(monkeypatch, api_env, tmp_path) -> None:
    from app.runtime.models import EmbedRequest, EmbedResult

    init()
    captured: list[EmbedRequest] = []

    def _embed(req: EmbedRequest) -> EmbedResult:
        captured.append(req)
        return EmbedResult(
            vectors=[[0.1, 0.2] for _ in req.texts],
            dimension=2,
            model=req.model,
        )

    monkeypatch.setattr("app.runtime.embeddings.embed", _embed)
    upsert_network(
        NetworkRow(
            id="net-out",
            name="kb",
            description=None,
            tags=[],
            document={"schemaVersion": 1, "name": "kb", "nodes": [], "edges": []},
            updated_at=utc_now(),
            last_used_at=None,
            last_run_id=None,
        )
    )
    folder = tmp_path / "stories"
    folder.mkdir()
    (folder / "note.md").write_text("# Hello\nworld", encoding="utf-8")
    data_dir = load_bootstrap().data_dir
    node = GraphNode(
        id="k-out",
        type="knowledge",
        position={"x": 0, "y": 0},
        data={
            "sourcePath": str(folder),
            "embeddingProvider": "ollama",
            "embeddingModel": "nomic-embed-text",
        },
    )
    assert index_node("net-out", node, data_dir=str(data_dir), force=True) == "ready"
    assert captured
    assert any("world" in text for text in captured[0].texts)


def test_index_reads_office_html_json_csv_and_code(monkeypatch, api_env, tmp_path) -> None:
    from docx import Document
    from openpyxl import Workbook

    from app.runtime.models import EmbedResult

    init()

    def _embed(req):
        return EmbedResult(vectors=[[0.1, 0.2] for _ in req.texts], dimension=2, model=req.model)

    monkeypatch.setattr("app.runtime.embeddings.embed", _embed)
    upsert_network(
        NetworkRow(
            id="net-files",
            name="kb",
            description=None,
            tags=[],
            document={"schemaVersion": 1, "name": "kb", "nodes": [], "edges": []},
            updated_at=utc_now(),
            last_used_at=None,
            last_run_id=None,
        )
    )
    folder = tmp_path / "kb"
    folder.mkdir()
    (folder / "skip.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    (folder / "hello.ts").write_text("export const city = 'Dortmund';", encoding="utf-8")
    (folder / "page.html").write_text("<h1>Intro</h1><p>Web text</p>", encoding="utf-8")
    (folder / "meta.json").write_text('{"role":"author"}', encoding="utf-8")
    (folder / "rows.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("Word body")
    doc.save(folder / "brief.docx")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "excel-cell"
    wb.save(folder / "sheet.xlsx")
    node = GraphNode(
        id="k-files",
        type="knowledge",
        position={"x": 0, "y": 0},
        data={
            "sourcePath": str(folder),
            "embeddingProvider": "ollama",
            "embeddingModel": "nomic-embed-text",
        },
    )
    assert index_node("net-files", node, data_dir=str(tmp_path / "appdata"), force=True) == "ready"
    from app.db.network_rag import list_chunks

    texts = " ".join(chunk.text for chunk in list_chunks("net-files", "k-files"))
    assert "Dortmund" in texts
    assert "Web text" in texts
    assert "author" in texts
    assert "Word body" in texts
    assert "excel-cell" in texts
    assert "PNG" not in texts


def _empty_network(network_id: str) -> None:
    upsert_network(
        NetworkRow(
            id=network_id,
            name="kb",
            description=None,
            tags=[],
            document={"schemaVersion": 1, "name": "kb", "nodes": [], "edges": []},
            updated_at=utc_now(),
            last_used_at=None,
            last_run_id=None,
        )
    )


def _knowledge_node(node_id: str, folder: Path) -> GraphNode:
    return GraphNode(
        id=node_id,
        type="knowledge",
        position={"x": 0, "y": 0},
        data={"sourcePath": str(folder), "embeddingModel": "nomic-embed-text"},
    )


def _embed_counter(monkeypatch):
    from app.runtime.models import EmbedResult

    calls = {"n": 0}

    def _embed(req):
        calls["n"] += 1
        return EmbedResult(vectors=[[0.1, 0.2] for _ in req.texts], dimension=2, model=req.model)

    monkeypatch.setattr("app.runtime.embeddings.embed", _embed)
    return calls


def test_edit_or_delete_rebuilds_the_whole_index(monkeypatch, api_env, tmp_path) -> None:
    from app.db.network_rag import get_collection, list_chunks

    init()
    calls = _embed_counter(monkeypatch)
    _empty_network("net-fp")
    folder = tmp_path / "notes"
    folder.mkdir()
    (folder / "keep.md").write_text("# Keep\nalpha", encoding="utf-8")
    (folder / "gone.md").write_text("# Gone\nbeta", encoding="utf-8")
    node = _knowledge_node("k-fp", folder)
    data_dir = str(tmp_path / "appdata")
    assert index_node("net-fp", node, data_dir=data_dir) == "ready"
    assert calls["n"] == 1
    first = get_collection("net-fp", "k-fp")
    assert first is not None and first.content_fingerprint

    (folder / "gone.md").unlink()
    (folder / "keep.md").write_text("# Keep\nalpha-two", encoding="utf-8")
    assert index_node("net-fp", node, data_dir=data_dir) == "ready"
    assert calls["n"] == 2
    texts = " ".join(chunk.text for chunk in list_chunks("net-fp", "k-fp"))
    assert "alpha-two" in texts
    assert "beta" not in texts
    second = get_collection("net-fp", "k-fp")
    assert second is not None
    assert second.content_fingerprint != first.content_fingerprint

    assert index_node("net-fp", node, data_dir=data_dir) == "ready"
    assert calls["n"] == 2


def test_missing_fingerprint_forces_reindex(monkeypatch, api_env, tmp_path) -> None:
    from app.db.network_rag import CollectionMeta, get_collection, upsert_collection

    init()
    calls = _embed_counter(monkeypatch)
    _empty_network("net-null")
    folder = tmp_path / "notes"
    folder.mkdir()
    (folder / "note.md").write_text("# Hello\nworld", encoding="utf-8")
    node = _knowledge_node("k-null", folder)
    data_dir = str(tmp_path / "appdata")
    assert index_node("net-null", node, data_dir=data_dir) == "ready"
    stored = get_collection("net-null", "k-null")
    assert stored is not None
    upsert_collection(
        CollectionMeta(
            network_id=stored.network_id,
            node_id=stored.node_id,
            source_path=stored.source_path,
            embedding_model_id=stored.embedding_model_id,
            dimension=stored.dimension,
            state="ready",
            updated_at=stored.updated_at,
            content_fingerprint=None,
        )
    )
    assert index_node("net-null", node, data_dir=data_dir) == "ready"
    assert calls["n"] == 2
    refreshed = get_collection("net-null", "k-null")
    assert refreshed is not None and refreshed.content_fingerprint


def test_vector_count_mismatch_marks_error_without_chunks(monkeypatch, api_env, tmp_path) -> None:
    from app.runtime.models import EmbedResult
    from app.db.network_rag import get_collection, list_chunks

    init()

    def _embed(req):
        return EmbedResult(vectors=[[0.1, 0.2]], dimension=2, model=req.model)

    monkeypatch.setattr("app.runtime.embeddings.embed", _embed)
    _empty_network("net-mismatch")
    folder = tmp_path / "notes"
    folder.mkdir()
    (folder / "a.md").write_text("# A\nalpha", encoding="utf-8")
    (folder / "b.md").write_text("# B\nbeta", encoding="utf-8")
    node = _knowledge_node("k-bad", folder)
    assert index_node("net-mismatch", node, data_dir=str(tmp_path / "appdata")) == "error"
    stored = get_collection("net-mismatch", "k-bad")
    assert stored is not None
    assert stored.state == "error"
    assert stored.content_fingerprint is None
    assert list_chunks("net-mismatch", "k-bad") == []


def test_index_skips_corpus_and_data_descendants(monkeypatch, api_env, tmp_path) -> None:
    from app.db.network_rag import list_chunks

    init()
    calls = _embed_counter(monkeypatch)
    _empty_network("net-skip")
    folder = tmp_path / "library"
    folder.mkdir()
    (folder / "note.md").write_text("# Public\npublic-note", encoding="utf-8")
    corpus = folder / "agentus_network_rag_documents"
    corpus.mkdir()
    (corpus / "secret.md").write_text("# Secret\ncorpus-secret", encoding="utf-8")
    data_dir = folder / "appdata"
    data_dir.mkdir()
    (data_dir / "private.md").write_text("# Private\ndata-secret", encoding="utf-8")
    node = _knowledge_node("k-skip", folder)
    assert index_node("net-skip", node, data_dir=str(data_dir)) == "ready"
    texts = " ".join(chunk.text for chunk in list_chunks("net-skip", "k-skip"))
    assert "public-note" in texts
    assert "corpus-secret" not in texts
    assert "data-secret" not in texts

    (corpus / "secret.md").write_text("# Secret\ncorpus-changed", encoding="utf-8")
    (data_dir / "private.md").write_text("# Private\ndata-changed", encoding="utf-8")
    assert index_node("net-skip", node, data_dir=str(data_dir)) == "ready"
    assert calls["n"] == 1


def test_index_rejects_a_data_dir_root(monkeypatch, api_env, tmp_path) -> None:
    from app.db.network_rag import get_collection

    init()
    calls = _embed_counter(monkeypatch)
    _empty_network("net-data")
    data_dir = tmp_path / "appdata"
    data_dir.mkdir()
    (data_dir / "note.md").write_text("# Hello\nworld", encoding="utf-8")
    node = _knowledge_node("k-data", data_dir)
    assert index_node("net-data", node, data_dir=str(data_dir)) == "error"
    assert calls["n"] == 0
    stored = get_collection("net-data", "k-data")
    assert stored is not None and stored.state == "error"
