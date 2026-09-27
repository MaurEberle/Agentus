from __future__ import annotations

import json
from pathlib import Path

from app.common.extract import extract_text, is_indexable


def test_plain_and_code(tmp_path: Path) -> None:
    md = tmp_path / "note.md"
    md.write_text("# Hello\nworld", encoding="utf-8")
    cs = tmp_path / "Greeter.cs"
    cs.write_text("namespace Demo { class Greeter {} }", encoding="utf-8")
    ts = tmp_path / "greeter.ts"
    ts.write_text("export const n = 1;", encoding="utf-8")
    png = tmp_path / "logo.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n")
    ole = tmp_path / "legacy.doc"
    ole.write_bytes(b"\xd0\xcf\x11\xe0")
    assert is_indexable(md)
    assert is_indexable(cs)
    assert is_indexable(ts)
    assert not is_indexable(png)
    assert not is_indexable(ole)
    assert "world" in extract_text(md)
    assert "namespace Demo" in extract_text(cs)
    assert extract_text(png) == ""


def test_html_json_csv(tmp_path: Path) -> None:
    html = tmp_path / "page.html"
    html.write_text(
        "<html><head><style>p{}</style></head><body><h1>Intro</h1>"
        "<p>Hello</p><script>alert(1)</script></body></html>",
        encoding="utf-8",
    )
    data = tmp_path / "data.json"
    data.write_text(json.dumps({"city": "Dortmund", "ok": True}), encoding="utf-8")
    table = tmp_path / "rows.csv"
    table.write_text("name;qty\napple;2\n", encoding="utf-8")
    html_text = extract_text(html)
    assert "Intro" in html_text
    assert "Hello" in html_text
    assert "alert" not in html_text
    assert "# Intro" in html_text
    json_text = extract_text(data)
    assert "Dortmund" in json_text
    csv_text = extract_text(table)
    assert "apple" in csv_text
    assert "2" in csv_text


def test_docx_xlsx_ipynb(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import Workbook

    word = tmp_path / "brief.docx"
    doc = Document()
    doc.add_heading("Titel", level=1)
    doc.add_paragraph("Der Absatz.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.save(word)

    excel = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Preise"
    ws["A1"] = "item"
    ws["B1"] = "eur"
    ws["A2"] = "apfel"
    ws["B2"] = 1.2
    wb.save(excel)

    notebook = tmp_path / "lab.ipynb"
    notebook.write_text(
        json.dumps(
            {
                "cells": [
                    {"cell_type": "markdown", "source": ["# Note\n", "body"]},
                    {"cell_type": "code", "source": "print(1)", "outputs": []},
                ]
            }
        ),
        encoding="utf-8",
    )

    word_text = extract_text(word)
    assert "# Titel" in word_text
    assert "Der Absatz." in word_text
    assert "A | B" in word_text
    excel_text = extract_text(excel)
    assert "Preise" in excel_text
    assert "apfel" in excel_text
    nb_text = extract_text(notebook)
    assert "body" in nb_text
    assert "print(1)" in nb_text
