"""Turn knowledge and help-corpus files into plain text for chunking."""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from html.parser import HTMLParser
from pathlib import Path

log = logging.getLogger("agentus.extract")

MAX_SOURCE_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 400_000

SKIP_SUFFIXES = frozenset(
    {
        ".sqlite",
        ".db",
        ".exe",
        ".dll",
        ".so",
        ".dylib",
        ".bin",
        ".png",
        ".jpg",
        ".jpeg",
        ".gif",
        ".webp",
        ".bmp",
        ".tif",
        ".tiff",
        ".ico",
        ".heic",
        ".mp3",
        ".wav",
        ".ogg",
        ".flac",
        ".mp4",
        ".webm",
        ".avi",
        ".mov",
        ".zip",
        ".7z",
        ".rar",
        ".gz",
        ".tar",
        ".doc",
        ".xls",
        ".ppt",
        ".pptx",
    }
)

PLAIN_SUFFIXES = frozenset({".md", ".markdown", ".mdx", ".txt", ".rst"})
HTML_SUFFIXES = frozenset({".html", ".htm", ".xhtml"})
JSON_SUFFIXES = frozenset({".json", ".jsonl"})
CSV_SUFFIXES = frozenset({".csv", ".tsv"})
PDF_SUFFIXES = frozenset({".pdf"})
DOCX_SUFFIXES = frozenset({".docx"})
XLSX_SUFFIXES = frozenset({".xlsx", ".xlsm"})
NOTEBOOK_SUFFIXES = frozenset({".ipynb"})
CODE_SUFFIXES = frozenset(
    {
        ".cs",
        ".csx",
        ".fs",
        ".vb",
        ".csproj",
        ".fsproj",
        ".vbproj",
        ".sln",
        ".ts",
        ".tsx",
        ".js",
        ".jsx",
        ".mjs",
        ".cjs",
        ".mts",
        ".cts",
        ".py",
        ".pyi",
        ".pyw",
        ".java",
        ".kt",
        ".kts",
        ".scala",
        ".groovy",
        ".go",
        ".rs",
        ".swift",
        ".c",
        ".h",
        ".cpp",
        ".cc",
        ".cxx",
        ".hpp",
        ".hh",
        ".m",
        ".mm",
        ".rb",
        ".php",
        ".lua",
        ".pl",
        ".pm",
        ".r",
        ".jl",
        ".sh",
        ".bash",
        ".zsh",
        ".ps1",
        ".psm1",
        ".bat",
        ".cmd",
        ".sql",
        ".graphql",
        ".gql",
        ".css",
        ".scss",
        ".less",
        ".sass",
        ".vue",
        ".svelte",
        ".astro",
        ".xml",
        ".xsl",
        ".xsd",
        ".svg",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
        ".conf",
        ".properties",
        ".gradle",
        ".cmake",
        ".mk",
        ".proto",
        ".thrift",
        ".dart",
        ".ex",
        ".exs",
        ".erl",
        ".hs",
        ".clj",
        ".cljs",
        ".tf",
        ".hcl",
    }
)


def _optional(name: str) -> bool:
    try:
        __import__(name)
    except ImportError:
        return False
    return True


def indexable_suffixes() -> frozenset[str]:
    out = set(
        PLAIN_SUFFIXES
        | HTML_SUFFIXES
        | JSON_SUFFIXES
        | CSV_SUFFIXES
        | CODE_SUFFIXES
        | NOTEBOOK_SUFFIXES
    )
    if _optional("pypdf"):
        out |= PDF_SUFFIXES
    if _optional("docx"):
        out |= DOCX_SUFFIXES
    if _optional("openpyxl"):
        out |= XLSX_SUFFIXES
    return frozenset(out)


def is_indexable(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix in SKIP_SUFFIXES:
        return False
    return suffix in indexable_suffixes()


def extract_text(path: Path) -> str:
    if not is_indexable(path):
        return ""
    try:
        size = path.stat().st_size
    except OSError:
        return ""
    if size > MAX_SOURCE_BYTES:
        log.warning("extract.skip.size path=%s size=%s", path, size)
        return ""
    suffix = path.suffix.lower()
    try:
        raw = _dispatch(path, suffix)
    except Exception:
        log.exception("extract.failed path=%s", path)
        return ""
    text = (raw or "").strip()
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
    return text


def _dispatch(path: Path, suffix: str) -> str:
    if suffix in PDF_SUFFIXES:
        return _read_pdf(path)
    if suffix in DOCX_SUFFIXES:
        return _read_docx(path)
    if suffix in XLSX_SUFFIXES:
        return _read_xlsx(path)
    if suffix in HTML_SUFFIXES:
        return _read_html(path)
    if suffix in NOTEBOOK_SUFFIXES:
        return _read_ipynb(path)
    if suffix in JSON_SUFFIXES:
        return _read_json(path)
    if suffix in CSV_SUFFIXES:
        return _read_csv(path)
    return _read_text(path)


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    if b"\x00" in data[:8192]:
        return ""
    return data.decode("utf-8-sig", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _heading_level(style_name: str, style_id: str) -> int | None:
    blob = f"{style_name} {style_id}".lower()
    for level in range(1, 7):
        if f"heading {level}" in blob or f"heading{level}" in blob:
            return level
        if f"überschrift {level}" in blob or f"überschrift{level}" in blob:
            return level
    return None


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = paragraph.style
        name = style.name if style is not None else ""
        style_id = getattr(style, "style_id", "") if style is not None else ""
        level = _heading_level(name or "", style_id or "")
        if level:
            parts.append("#" * level + " " + text)
        else:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        blocks: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                if any(cells):
                    rows.append(" | ".join(cells))
                if len(rows) >= 20_000:
                    break
            if rows:
                title = sheet.title.strip() or "sheet"
                blocks.append(f"# {title}\n" + "\n".join(rows))
        return "\n\n".join(blocks)
    finally:
        workbook.close()


class _HtmlText(HTMLParser):
    _BLOCK = frozenset(
        {
            "p",
            "div",
            "section",
            "article",
            "header",
            "footer",
            "li",
            "tr",
            "br",
            "hr",
            "table",
            "ul",
            "ol",
            "blockquote",
            "pre",
        }
    )
    _SKIP = frozenset({"script", "style", "noscript", "template"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag in self._SKIP:
            self._skip += 1
            return
        if self._skip:
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._parts.append("\n" + "#" * int(tag[1]) + " ")
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self._SKIP and self._skip:
            self._skip -= 1
            return
        if self._skip:
            return
        if tag in self._BLOCK or tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip and data:
            self._parts.append(data)

    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", "".join(self._parts)).strip()


def _read_html(path: Path) -> str:
    parser = _HtmlText()
    parser.feed(_read_text(path))
    parser.close()
    return parser.text()


def _pretty_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


def _read_json(path: Path) -> str:
    raw = _read_text(path)
    if path.suffix.lower() == ".jsonl":
        blocks: list[str] = []
        for line in raw.splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                blocks.append(_pretty_json(json.loads(line)))
            except json.JSONDecodeError:
                blocks.append(line)
        return "\n\n".join(blocks)
    try:
        return _pretty_json(json.loads(raw))
    except json.JSONDecodeError:
        return raw


def _read_ipynb(path: Path) -> str:
    try:
        notebook = json.loads(_read_text(path))
    except json.JSONDecodeError:
        return ""
    parts: list[str] = []
    cells = notebook.get("cells") if isinstance(notebook, dict) else None
    if not isinstance(cells, list):
        return ""
    for index, cell in enumerate(cells, 1):
        if not isinstance(cell, dict):
            continue
        source = cell.get("source")
        if isinstance(source, list):
            text = "".join(str(part) for part in source)
        else:
            text = str(source or "")
        text = text.strip()
        if not text:
            continue
        kind = str(cell.get("cell_type") or "cell")
        parts.append(f"# {kind} {index}\n{text}")
    return "\n\n".join(parts)


def _read_csv(path: Path) -> str:
    raw = _read_text(path)
    if not raw.strip():
        return ""
    dialect: csv.Dialect | type[csv.Dialect] = csv.excel
    if path.suffix.lower() == ".tsv":
        dialect = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
        except csv.Error:
            pass
    rows: list[str] = []
    reader = csv.reader(io.StringIO(raw), dialect)
    for row in reader:
        cells = [cell.strip() for cell in row]
        if any(cells):
            rows.append(" | ".join(cells))
        if len(rows) >= 20_000:
            break
    return "\n".join(rows)
